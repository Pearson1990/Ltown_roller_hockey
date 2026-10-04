import openpyxl, json, re, sys, datetime, warnings
warnings.filterwarnings('ignore')
src=sys.argv[1]
wb=openpyxl.load_workbook(src,data_only=True)
# rosters
ws=wb['Winter 2026-2027 Rosters']; rows=list(ws.iter_rows(min_row=2,max_row=11,min_col=2,max_col=10,values_only=True))
teams={}
for j,t in enumerate(rows[0]):
    name=t.replace("'s Team",'')
    teams[name]=[r[j] for r in rows[1:] if r[j]]
# schedule
ws=wb['Winter 2026-27 Schedule-Results']
games=[];byes={};dates={};wk=None
for r in ws.iter_rows(min_row=9,max_row=160,max_col=8,values_only=True):
    a,b,c,d,e,f,g,_=r
    if a and str(a).startswith('Week'): wk=int(a.split()[1]); dates[wk]=b
    if not d or wk is None: continue
    if d.startswith('Bye'):
        n=d.replace('Bye:','').strip(); 
        if n: byes[wk]=n.replace("'s","")
        continue
    m=re.match(r"(.+?) vs\. (.+)",d)
    if not m: continue
    h,aw=[x.strip().replace("'s","") for x in m.groups()]
    games.append(dict(week=wk,date=dates.get(wk),time=c.strftime('%H:%M') if c else None,home=h,away=aw,hs=e,as_=f,refs=g))
# players
ws=wb['Winter 2026-2027 Player Stats']
players=[];goalies=[]
hdr=[c.value for c in ws[1]]
nweeks=sum(1 for v in hdr if v and str(v).startswith('Week'))
for r in ws.iter_rows(min_row=4,max_row=105,values_only=True):
    n=r[0]
    if not n or 'Sub Goalie' in n or not re.search(r"\(([^()]+?)'s Team\)$",n): continue
    team=re.search(r"\(([^()]+?)'s Team\)$",n).group(1); nm=re.sub(r" \([^()]+?'s Team\)$","",n)
    wkly=[]
    for w in range(nweeks):
        b=7+5*w; blk=r[b:b+5]
        wkly.append([x or 0 for x in blk[:4]]+[1 if blk[4] else 0] if len(blk)==5 else [0]*5)
    if '(G)' in nm: continue
    players.append(dict(name=nm,team=team,g=r[1] or 0,a=r[2] or 0,pim=r[3] or 0,pts=r[4] or 0,gp=r[6] or 0,weekly=wkly))
# goalies from stats sorted
ws=wb['Winter 2026-2027 Stats Sorted']
for r in ws.iter_rows(min_row=15,max_row=23,max_col=9,values_only=True):
    m=re.match(r"(.+?) \(G\) \((.+?)'s Team\)",r[1])
    if m: goalies.append(dict(name=m.group(1),team=m.group(2),gp=r[2] or 0,w=r[3] or 0,l=r[4] or 0,t=r[5] or 0,ga=r[6] or 0))
    elif 'Sub (G)' in r[1]: goalies.append(dict(name='Sub',team='Billy',gp=0,w=0,l=0,t=0,ga=0))
# all-time
ws=wb['All Time Total Stats']; alltime=[]
for r in ws.iter_rows(min_row=6,max_row=283,max_col=6,values_only=True):
    if r[0] and r[0] not in('TBD','(blank)') and (r[1] or 0)>0: alltime.append(dict(name=r[0],gp=r[1],g=r[2],a=r[3],pts=r[4],pim=r[5]))
json.dump(dict(teams=teams,games=games,byes=byes,players=players,goalies=goalies,alltime=alltime,asof='2026-09-30'),open('data.json','w'),separators=(',',':'))
print(len(players),len(goalies),len(games),len(alltime),nweeks)
